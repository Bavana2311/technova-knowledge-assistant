import os
from typing import List

from langchain_core.documents import Document

import requests
import pypdf
import docx


#step-1:Convert the plain text into a list of langchain documents
def text_to_doc(text: str, source: str = "inline") -> List[Document]:
    return [Document(page_content=text, metadata={"source": source})]

#Step-2: Load the .txt, .pdf, .docx, and  convert them into a list of langchain documents
def load_txt(path: str) -> List[Document]:
    if not os.path.exists(path):
        raise FileNotFoundError("File not found, Provide a valid file path.")
    with open(path, "r", encoding="utf-8") as f:
        text=f.read()
    return text_to_doc(text, source=os.path.abspath(path))

def load_pdf(path: str) -> List[Document]:
    reader = pypdf.PdfReader(path)
    docs= []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = text.replace("\n", " ")
        text = " ".join(text.split())
        docs.extend(text_to_doc(text, source= os.path.abspath(path)))
    return docs

def load_docx(path: str) -> List[Document]:
    doc = docx.Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    text = "\n\n".join(paragraphs)
    return text_to_doc(text, source = os.path.abspath(path))


def load_documents(path_or_url: str) -> List[Document]:
    
    ext = os.path.splitext(path_or_url)[1].lower()
    if ext in [".txt", ".md"]:
        return load_txt(path_or_url)
    elif ext == ".pdf":
        return load_pdf(path_or_url)
    elif ext in [".docx", ".docs", ".doc"]:
        return load_docx(path_or_url)
    else:
        try:
            return load_txt(path_or_url)
        except Exception as e:
            raise ValueError(f"Unsupported file type: {ext}. Error: {e}")
        